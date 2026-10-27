# -*- coding: utf-8 -*-
"""Gera .xlsx no template GherkinFlow (7 colunas + Foursys Dark)."""
from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "222239"
NAVY2 = "2E2E4A"
ORANGE = "FF5315"
MINT = "89BAB1"
VANILLA = "FFE2A9"
OFF = "F7F6F2"
WHITE = "FFFFFF"
MUTED = "5C5C72"

HEADERS_7 = [
    "#",
    "Funcionalidade (Feature)",
    "Dado (Given)",
    "Quando (When)",
    "Então (Then)",
    "E… (And)",
    "Referência do Código",
]

HEADERS_TODOS = [
    "ID",
    "Tela",
    "Tipo",
    "Funcionalidade (Feature)",
    "Dado (Given)",
    "Quando (When)",
    "Então (Then)",
    "E… (And)",
    "Referência do Código",
    "Status E2E",
    "Duração (ms)",
]


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def thin() -> Border:
    s = Side(style="thin", color="E8E6E0")
    return Border(left=s, right=s, top=s, bottom=s)


def header_font() -> Font:
    return Font(name="Calibri", bold=True, color=WHITE, size=11)


def body_font(bold=False, color="181828", size=10) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size)


def wrap() -> Alignment:
    return Alignment(wrap_text=True, vertical="top")


def parse_bdd_md(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    rows = []
    tela = ""
    for line in text.splitlines():
        if line.startswith("## "):
            tela = line[3:].strip()
            continue
        if not line.startswith("| AUTH-"):
            continue
        parts = [p.strip() for p in line.strip("|").split("|")]
        if len(parts) < 7:
            continue
        feat = parts[1]
        tipo = "Positivo"
        if "[Negativo]" in feat:
            tipo = "Negativo"
        elif "[Regressivo]" in feat:
            tipo = "Regressivo"
        titulo = re.sub(r"^\[(Positivo|Negativo|Regressivo)\]\s*", "", feat).strip()
        rows.append(
            {
                "id": parts[0],
                "tela": tela,
                "tipo": tipo,
                "feature": f"{tipo} · {titulo}",
                "given": parts[2],
                "when": parts[3],
                "then": parts[4],
                "and": parts[5],
                "ref": parts[6].replace("`", ""),
            }
        )
    return rows


def flatten_specs(suites, acc=None):
    acc = acc if acc is not None else []
    for suite in suites or []:
        for spec in suite.get("specs") or []:
            acc.append(spec)
        flatten_specs(suite.get("suites"), acc)
    return acc


def spec_status(spec: dict) -> str:
    tests = spec.get("tests") or []
    results = [r for t in tests for r in (t.get("results") or [])]
    if any(r.get("status") in ("failed", "timedOut", "unexpected") for r in results):
        return "failed"
    if tests and all(
        t.get("expectedStatus") == "skipped"
        or all(r.get("status") == "skipped" for r in (t.get("results") or []))
        for t in tests
    ):
        return "skipped"
    if spec.get("ok"):
        return "passed"
    return (results[0].get("status") if results else "unknown") or "unknown"


def spec_duration(spec: dict) -> int:
    tests = spec.get("tests") or []
    results = [r for t in tests for r in (t.get("results") or [])]
    return int(sum(r.get("duration") or 0 for r in results))


def load_exec(json_path: Path) -> dict[str, dict]:
    if not json_path.exists():
        return {}
    data = json.loads(json_path.read_text(encoding="utf-8"))
    out = {}
    for spec in flatten_specs(data.get("suites")):
        title = spec.get("title") or ""
        m = re.match(r"^([A-Z]+-\d+)", title)
        key = m.group(1) if m else title[:24]
        out[key] = {"status": spec_status(spec), "duration": spec_duration(spec), "title": title}
    return out


def style_header(ws, cols: int):
    for col in range(1, cols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill(NAVY)
        cell.font = header_font()
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=ORANGE))


def style_body(ws, start_row: int, end_row: int, cols: int, tipo_col: int | None = None):
    for r in range(start_row, end_row + 1):
        even = r % 2 == 0
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.font = body_font()
            cell.alignment = wrap()
            cell.border = thin()
            cell.fill = fill("F3F8F6") if even else fill(WHITE)
        if tipo_col:
            tipo = str(ws.cell(r, tipo_col).value or "")
            if "Positivo" in tipo:
                ws.cell(r, tipo_col).fill = fill(MINT)
                ws.cell(r, tipo_col).font = body_font(bold=True, color=NAVY)
            elif "Negativo" in tipo:
                ws.cell(r, tipo_col).fill = fill(VANILLA)
                ws.cell(r, tipo_col).font = body_font(bold=True, color=NAVY)
            elif "Regressivo" in tipo:
                ws.cell(r, tipo_col).fill = fill(NAVY2)
                ws.cell(r, tipo_col).font = body_font(bold=True, color=WHITE)


def autosize(ws, widths: list[int], *, filter_on: bool = True, freeze: str = "A2"):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = freeze
    if filter_on:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_meta(ws, meta: dict):
    ws["A1"] = "GherkinFlow · Tabela BDD (template Foursys)"
    ws["A1"].font = Font(name="Calibri", bold=True, color=ORANGE, size=16)
    ws.merge_cells("A1:B1")
    rows = [
        ("Projeto", meta["projeto"]),
        ("Tela / módulo", meta["modulo"]),
        ("Versão", meta["versao"]),
        ("Gerado", meta["gerado"]),
        ("Agente", "GherkinFlow + Playwright (esteira QA)"),
        ("Total", meta["total"]),
        ("Positivo", meta["pos"]),
        ("Negativo", meta["neg"]),
        ("Regressivo", meta["reg"]),
        ("E2E passed", meta["passed"]),
        ("E2E failed", meta["failed"]),
        ("E2E skipped", meta["skipped"]),
        ("Fonte HTML", meta["fonte"]),
        ("Par MD", meta["md"]),
    ]
    ws["A2"] = "Meta"
    ws["B2"] = "Valor"
    for col in (1, 2):
        cell = ws.cell(2, col)
        cell.fill = fill(NAVY)
        cell.font = header_font()
        cell.border = Border(bottom=Side(style="medium", color=ORANGE))
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(i, 1, k).font = body_font(bold=True, color=NAVY)
        ws.cell(i, 2, v).font = body_font()
        ws.cell(i, 1).fill = fill(OFF)
        ws.cell(i, 2).fill = fill(WHITE)
        ws.cell(i, 1).border = thin()
        ws.cell(i, 2).border = thin()
    autosize(ws, [28, 80], filter_on=False, freeze="A3")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bdd", required=True)
    ap.add_argument("--json", default="")
    ap.add_argument("--out", required=True)
    ap.add_argument("--modulo", default="ep05-auth")
    ap.add_argument("--nome", default="PRISMA-EP-05-AUTH")
    ap.add_argument("--version", default="1")
    args = ap.parse_args()

    bdd_path = Path(args.bdd)
    rows = parse_bdd_md(bdd_path)
    exec_map = load_exec(Path(args.json)) if args.json else {}

    pos = sum(1 for r in rows if r["tipo"] == "Positivo")
    neg = sum(1 for r in rows if r["tipo"] == "Negativo")
    reg = sum(1 for r in rows if r["tipo"] == "Regressivo")
    passed = sum(1 for r in rows if exec_map.get(r["id"], {}).get("status") == "passed")
    failed = sum(1 for r in rows if exec_map.get(r["id"], {}).get("status") == "failed")
    skipped = sum(1 for r in rows if exec_map.get(r["id"], {}).get("status") == "skipped")

    wb = Workbook()

    ws_meta = wb.active
    ws_meta.title = "Meta"
    write_meta(
        ws_meta,
        {
            "projeto": "Prisma / Portal de Contestação",
            "modulo": args.nome,
            "versao": f"v{args.version}-{date.today().strftime('%Y%m%d')}",
            "gerado": date.today().isoformat(),
            "total": f"{len(rows)} ({pos} Positivo · {neg} Negativo · {reg} Regressivo)",
            "pos": pos,
            "neg": neg,
            "reg": reg,
            "passed": passed,
            "failed": failed,
            "skipped": skipped,
            "fonte": "11_BDD_PRISMA-EP-05-AUTH_PortalContestacao.html",
            "md": str(bdd_path).replace("\\", "/"),
        },
    )

    ws7 = wb.create_sheet("Tabela BDD")
    ws7.append(HEADERS_7)
    for r in rows:
        ws7.append([r["id"], r["feature"], r["given"], r["when"], r["then"], r["and"], r["ref"]])
    last = 1 + len(rows)
    style_header(ws7, 7)
    style_body(ws7, 2, last, 7)
    for i, r in enumerate(rows, 2):
        feat = ws7.cell(i, 2)
        if r["tipo"] == "Positivo":
            feat.fill = fill(MINT)
        elif r["tipo"] == "Negativo":
            feat.fill = fill(VANILLA)
        else:
            feat.fill = fill(NAVY2)
            feat.font = body_font(bold=True, color=WHITE)
    autosize(ws7, [12, 36, 28, 28, 28, 28, 40])

    ws_all = wb.create_sheet("Todos")
    ws_all.append(HEADERS_TODOS)
    for r in rows:
        ex = exec_map.get(r["id"], {})
        ws_all.append(
            [
                r["id"],
                r["tela"],
                r["tipo"],
                r["feature"],
                r["given"],
                r["when"],
                r["then"],
                r["and"],
                r["ref"],
                ex.get("status", ""),
                ex.get("duration", ""),
            ]
        )
    last_all = 1 + len(rows)
    style_header(ws_all, 11)
    style_body(ws_all, 2, last_all, 11, tipo_col=3)
    for i in range(2, last_all + 1):
        st = str(ws_all.cell(i, 10).value or "")
        cell = ws_all.cell(i, 10)
        if st == "passed":
            cell.fill = fill(MINT)
        elif st == "failed":
            cell.fill = fill(ORANGE)
            cell.font = body_font(bold=True, color=WHITE)
        elif st == "skipped":
            cell.fill = fill(VANILLA)
    autosize(ws_all, [12, 36, 14, 36, 28, 28, 28, 28, 40, 14, 14])

    telas = sorted({r["tela"] for r in rows})
    for tela in telas:
        name = "SCR-01" if "SCR-01" in tela else "SCR-02" if "SCR-02" in tela else tela[:31]
        ws = wb.create_sheet(name)
        ws.append(HEADERS_7)
        subset = [r for r in rows if r["tela"] == tela]
        for r in subset:
            ws.append([r["id"], r["feature"], r["given"], r["when"], r["then"], r["and"], r["ref"]])
        last_s = 1 + len(subset)
        style_header(ws, 7)
        style_body(ws, 2, last_s, 7)
        autosize(ws, [12, 36, 28, 28, 28, 28, 40])

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"OK {out} · {len(rows)} cenários")


if __name__ == "__main__":
    main()
